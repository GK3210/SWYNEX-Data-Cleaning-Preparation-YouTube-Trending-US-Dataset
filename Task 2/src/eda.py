"""
SWYNEX Task 2 - Exploratory Data Analysis
Input : data/youtube_trending_us_cleaned.csv  (output of Task 1)
Run   : python src/eda.py   (from the repository root)
Output: charts/*.png, reports/*.csv, reports/EDA_summary_tables.xlsx
"""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mt
import numpy as np
import pandas as pd
import seaborn as sns

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent if HERE.name == "src" else HERE   # works from src/ or from the repo root
CH, RP = ROOT / "charts", ROOT / "reports"
CH.mkdir(exist_ok=True)
RP.mkdir(exist_ok=True)
NAVY, TEAL, ORANGE, GREY = "#1F3864", "#2A9D8F", "#E76F51", "#9AA5B1"
sns.set_theme(style="whitegrid", font="DejaVu Sans")
plt.rcParams.update({"axes.titlesize": 11.5, "axes.titleweight": "bold", "axes.labelsize": 10,
                     "figure.dpi": 110, "figure.facecolor": "white", "savefig.facecolor": "white"})
M = mt.FuncFormatter(lambda x, _: f"{x/1e6:g}M" if x >= 1e6 else f"{x/1e3:.0f}K" if x >= 1e3 else f"{x:.0f}")


import textwrap
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import Rectangle

PDF = PdfPages(RP / "EDA_Figures_Report.pdf")
SOURCE = ("Source: YouTube Trending Videos (US), cleaned dataset from SWYNEX Task 1  |  14 Nov 2017 - 14 Jun 2018  |  "
          "n = 40,899 video-days, 6,351 unique videos")


def save(fig, name, number, title, note):
    """Report-ready figure: numbered title, note, source line and a complete border."""
    w = fig.get_size_inches()[0]
    fig.tight_layout(rect=(0.02, 0.105, 0.98, 0.925))
    fig.suptitle(f"Figure {number}: {title}", fontsize=14.5, fontweight="bold", color=NAVY, x=0.02, ha="left", y=0.975)
    fig.text(0.02, 0.062, textwrap.fill("Note: " + note, width=int(w * 14)), fontsize=8.5, style="italic",
             color="#333333", va="center", ha="left")
    fig.text(0.02, 0.018, SOURCE, fontsize=8, color="#555555", va="center", ha="left")
    fig.add_artist(Rectangle((0.003, 0.003), 0.994, 0.994, transform=fig.transFigure, fill=False,
                             edgecolor=NAVY, linewidth=3.5))
    fig.add_artist(Rectangle((0.0, 0.0), 1, 1, transform=fig.transFigure, fill=False, edgecolor="white", linewidth=0))
    fig.savefig(CH / name, dpi=200)
    PDF.savefig(fig)
    plt.close(fig)


def bar_labels(ax, fmt, horizontal=False, size=8, pad=3):
    for cont in ax.containers:
        ax.bar_label(cont, labels=[fmt(v) for v in cont.datavalues], padding=pad, fontsize=size)
    if horizontal:
        ax.margins(x=0.14)
    else:
        ax.margins(y=0.12)


def kfmt(v):
    return f"{v/1e6:.2f}M" if v >= 1e6 else f"{v/1e3:.0f}K"


df = pd.read_csv(ROOT / "data/youtube_trending_us_cleaned.csv", parse_dates=["publish_time", "trending_date"])
df["month"] = df["trending_date"].dt.to_period("M").astype(str)
DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

# ============================================================ 1. Overview & descriptive statistics
num = ["views", "likes", "dislikes", "comment_count", "days_to_trend", "tag_count", "title_length"]
desc = df[num].describe(percentiles=[.25, .5, .75, .95]).T
desc["skewness"] = df[num].skew()
desc["missing"] = df[num].isna().sum()
desc.round(2).to_csv(RP / "descriptive_statistics.csv")
overview = {
    "Rows": len(df), "Unique videos": df.video_id.nunique(), "Unique channels": df.channel_title.nunique(),
    "Categories": df.category_name.nunique(), "First trending date": df.trending_date.min().date(),
    "Last trending date": df.trending_date.max().date(),
    "Trending days": df.trending_date.nunique(),
    "Avg rows per trending day": round(len(df) / df.trending_date.nunique(), 1),
}
pd.Series(overview).to_csv(RP / "dataset_overview.csv", header=["value"])

# ============================================================ 2. Views distribution & outliers
q1, q3 = df.views.quantile([.25, .75])
iqr_hi = q3 + 1.5 * (q3 - q1)
n_out = int((df.views > iqr_hi).sum())
fig, ax = plt.subplots(1, 2, figsize=(15, 6.2), gridspec_kw={"width_ratios": [1, 1.15]})
ax[0].hist(np.log10(df.views), bins=60, color=NAVY, alpha=.92, edgecolor="white", linewidth=.3)
for v, c, l in [(df.views.median(), ORANGE, f"Median = {df.views.median():,.0f}"),
                (df.views.mean(), TEAL, f"Mean = {df.views.mean():,.0f}")]:
    ax[0].axvline(np.log10(v), color=c, ls="--", lw=2.2, label=l)
ax[0].set(title="(a) Distribution of views per video-day (log scale)", xlabel="Views (log scale)",
          ylabel="Number of video-days")
ax[0].set_xticks([3, 4, 5, 6, 7, 8]); ax[0].set_xticklabels(["1K", "10K", "100K", "1M", "10M", "100M"])
ax[0].legend(title="Statistic", loc="upper right")
order = df.groupby("category_name").views.median().sort_values(ascending=False).index
sns.boxplot(data=df, x="category_name", y="views", ax=ax[1], color="#A9C1E8", fliersize=1.5, order=order)
ax[1].set_yscale("log"); ax[1].yaxis.set_major_formatter(M)
ax[1].tick_params(axis="x", rotation=60); plt.setp(ax[1].get_xticklabels(), ha="right")
ax[1].set(title="(b) Views by category (box = middle 50%, dots = outliers)", xlabel="Category", ylabel="Views (log scale)")
save(fig, "01_views_distribution.png", 1, "Views are heavily right-skewed",
     f"The mean is {df.views.mean()/df.views.median():.1f}x the median. By the IQR rule, {n_out:,} rows "
     f"({n_out/len(df)*100:.1f}%) are outliers (above {iqr_hi/1e6:.2f}M views). Maximum = {df.views.max():,} views "
     "(Childish Gambino - This Is America).")

# ============================================================ 3. Category analysis
vid = df.groupby("video_id").agg(days_on_list=("trending_date", "nunique"), category_name=("category_name", "first"))
cat = df.groupby("category_name").agg(
    video_days=("video_id", "size"), unique_videos=("video_id", "nunique"),
    median_views=("views", "median"), mean_views=("views", "mean"),
    median_like_ratio=("like_ratio", "median"), median_dislike_ratio=("dislike_ratio", "median"),
    median_comments=("comment_count", "median"))
cat["share_of_rows_pct"] = (cat.video_days / len(df) * 100).round(1)
cat["avg_days_on_trending"] = vid.groupby("category_name").days_on_list.mean().round(2)
cat = cat.sort_values("video_days", ascending=False)
cat.round(4).to_csv(RP / "category_summary.csv")

fig, ax = plt.subplots(1, 2, figsize=(15, 6.6))
c1 = cat.sort_values("video_days")
ax[0].barh(c1.index, c1.video_days, color=NAVY)
for i, (v, p) in enumerate(zip(c1.video_days, c1.share_of_rows_pct)):
    ax[0].text(v + 90, i, f"{v:,} ({p}%)", va="center", fontsize=8)
ax[0].margins(x=0.22)
ax[0].set(title="(a) Volume: video-days on the trending list", xlabel="Video-days (count and % of all rows)", ylabel="Category")
c2 = cat.sort_values("median_views")
ax[1].barh(c2.index, c2.median_views, color=[ORANGE if x in ("Music", "Gaming", "Film & Animation") else GREY for x in c2.index])
for i, v in enumerate(c2.median_views):
    ax[1].text(v + 15000, i, kfmt(v), va="center", fontsize=8)
ax[1].margins(x=0.14); ax[1].xaxis.set_major_formatter(M)
ax[1].set(title="(b) Reach: median views per video-day", xlabel="Median views", ylabel="Category")
save(fig, "02_category_volume_vs_views.png", 2, "Entertainment & Music dominate volume; Music, Gaming & Film dominate reach",
     "Orange bars = the three highest-reach categories. Nonprofits & Activism and Shows have only 57 rows each, "
     "so their figures should be read with caution.")

fig, ax = plt.subplots(1, 2, figsize=(15, 6.6))
c3 = cat.sort_values("median_like_ratio")
ax[0].barh(c3.index, c3.median_like_ratio * 100, color=[ORANGE if x in ("News & Politics", "Sports") else TEAL for x in c3.index])
for i, v in enumerate(c3.median_like_ratio * 100):
    ax[0].text(v + .05, i, f"{v:.2f}%", va="center", fontsize=8)
ax[0].margins(x=0.14)
ax[0].set(title="(a) Median like ratio (likes per 100 views)", xlabel="Likes per 100 views (%)", ylabel="Category")
c4 = cat.sort_values("avg_days_on_trending")
ax[1].barh(c4.index, c4.avg_days_on_trending, color=[ORANGE if x in ("News & Politics", "Sports") else NAVY for x in c4.index])
for i, v in enumerate(c4.avg_days_on_trending):
    ax[1].text(v + .15, i, f"{v:.1f}", va="center", fontsize=8)
ax[1].margins(x=0.1)
ax[1].set(title="(b) Average days a video stays on the list", xlabel="Days on trending list", ylabel="Category")
save(fig, "03_category_engagement_and_lifespan.png", 3, "News & Sports get fewer likes and a shorter time on the list",
     "Orange bars highlight News & Politics and Sports. Like ratio = likes / views for rows where ratings are enabled. "
     "'Shows' (4 videos) stays longest on average but is a very small sample.")

# ============================================================ 4. Publish timing
heat = df.pivot_table(index="publish_day", columns="publish_hour", values="video_id", aggfunc="size").reindex(DAYS)
fig, ax = plt.subplots(figsize=(16, 6.2))
sns.heatmap(heat, cmap="Blues", ax=ax, annot=True, fmt="d", annot_kws={"size": 6.5}, linewidths=.4, linecolor="white",
            cbar_kws={"label": "Number of video-days"})
ax.set(title="Number of trending video-days by publish day and hour", xlabel="Publish hour (UTC, 0-23)", ylabel="Publish day")
ax.tick_params(axis="y", rotation=0)
save(fig, "04_publish_time_heatmap.png", 4, "Weekday afternoons (UTC) are the peak publishing window",
     f"{df.publish_hour.between(14, 17).mean()*100:.1f}% of trending rows were published between 14:00 and 17:00 UTC "
     "(about 9am-12pm US Eastern). Darker cells = more videos.")

wk = df.assign(weekend=df.publish_day.isin(["Saturday", "Sunday"]))
wk_share = wk.weekend.mean() * 100
wk_med = wk.groupby("weekend").views.median()
day_tbl = df.groupby("publish_day").agg(video_days=("video_id", "size"), median_views=("views", "median")).reindex(DAYS)
day_tbl.to_csv(RP / "publish_day_summary.csv")
fig, ax = plt.subplots(1, 2, figsize=(14, 5.8))
cols_wd = [NAVY] * 5 + [ORANGE] * 2
ax[0].bar(day_tbl.index.str[:3], day_tbl.video_days, color=cols_wd)
bar_labels(ax[0], lambda v: f"{v:,.0f}")
ax[0].set(title="(a) Video-days by publish day", xlabel="Day of week published", ylabel="Video-days")
ax[1].bar(day_tbl.index.str[:3], day_tbl.median_views, color=cols_wd)
bar_labels(ax[1], kfmt)
ax[1].yaxis.set_major_formatter(M); ax[1].set(title="(b) Median views by publish day", xlabel="Day of week published", ylabel="Median views")
save(fig, "05_weekday_volume_vs_views.png", 5, "Weekend uploads are fewer but earn higher median views",
     f"Orange = Saturday and Sunday. Weekends are {wk_share:.1f}% of rows (28.6% if evenly spread). "
     f"Median views: weekend {wk_med[True]:,.0f} vs weekday {wk_med[False]:,.0f}.")

# ============================================================ 5. Trend over time (pattern shift)
# first trending date is excluded: every video is "new" on day 1 of the data, which would inflate November
daily_new = df.sort_values("trending_date").drop_duplicates("video_id").groupby("trending_date").size().iloc[1:]
mon = df.groupby("month").agg(median_views=("views", "median"), median_days_to_trend=("days_to_trend", "median"))
mon["new_videos_per_day"] = daily_new.groupby(daily_new.index.to_period("M").astype(str)).mean().round(1)
mon.round(2).to_csv(RP / "monthly_trend.csv")
fig, ax = plt.subplots(1, 3, figsize=(16, 5.6))
labels = [pd.Period(m).strftime("%b %y") for m in mon.index]
for a, col, t, c, yl, f in [(ax[0], "median_views", "(a) Median views per video-day", NAVY, "Median views", kfmt),
                            (ax[1], "median_days_to_trend", "(b) Median days from upload to trending", ORANGE, "Days", lambda v: f"{v:.1f}"),
                            (ax[2], "new_videos_per_day", "(c) New videos entering the list per day", TEAL, "Videos per day", lambda v: f"{v:.0f}")]:
    a.plot(labels, mon[col], marker="o", color=c, lw=2.5)
    for x, v in zip(labels, mon[col]):
        a.annotate(f(v), (x, v), textcoords="offset points", xytext=(0, 8), ha="center", fontsize=8)
    a.set(title=t, xlabel="Trending month", ylabel=yl); a.tick_params(axis="x", rotation=45); a.margins(y=0.15)
ax[0].yaxis.set_major_formatter(M)
save(fig, "06_monthly_trend_shift.png", 6, "A clear shift from March 2018: bigger videos, slower to trend, less turnover",
     "Jun 2018 covers only 14 days (to 14 Jun). The first trending day (14 Nov 2017) is excluded from panel (c) because all "
     "200 videos are 'new' on day 1. A change in YouTube's trending algorithm is a plausible cause but cannot be proven here.")

# ============================================================ 6. Correlations & content features
cols = ["views", "likes", "dislikes", "comment_count", "tag_count", "title_length", "title_caps_ratio", "days_to_trend"]
pretty = ["Views", "Likes", "Dislikes", "Comments", "Tag count", "Title length", "Title caps ratio", "Days to trend"]
corr = df[cols].corr(method="spearman")
corr.round(3).to_csv(RP / "correlation_spearman.csv")
fig, ax = plt.subplots(figsize=(9.5, 7.8))
sns.heatmap(corr.set_axis(pretty, axis=0).set_axis(pretty, axis=1), annot=True, fmt=".2f", cmap="RdBu_r", center=0,
            vmin=-1, vmax=1, ax=ax, square=True, linewidths=.5, cbar_kws={"label": "Spearman correlation (-1 to +1)"})
ax.set(title="Spearman rank correlation between numeric variables")
ax.tick_params(axis="x", rotation=45); plt.setp(ax.get_xticklabels(), ha="right")
save(fig, "07_correlation_heatmap.png", 7, "Engagement metrics rise together; tags and title style barely matter",
     f"Views-likes = {corr.loc['views','likes']:.2f}, views-comments = {corr.loc['views','comment_count']:.2f}, "
     f"views-tags = {corr.loc['views','tag_count']:.2f}. Likes, dislikes and comments exclude rows where ratings/comments are disabled. "
     "Correlation does not imply causation.")

df["title_len_band"] = pd.cut(df.title_length, [0, 30, 50, 70, 100], labels=["Up to 30", "31-50", "51-70", "71-100"])
df["tag_band"] = pd.cut(df.tag_count, [-1, 0, 10, 20, 30, 100], labels=["0", "1-10", "11-20", "21-30", "31+"])
tl = df.groupby("title_len_band", observed=True).agg(median_like_ratio=("like_ratio", "median"), median_views=("views", "median"))
tg = df.groupby("tag_band", observed=True).agg(median_views=("views", "median"), video_days=("views", "size"))
tl.round(4).to_csv(RP / "title_length_summary.csv"); tg.to_csv(RP / "tag_count_summary.csv")
fig, ax = plt.subplots(1, 2, figsize=(14, 5.8))
ax[0].bar(tl.index.astype(str), tl.median_like_ratio * 100, color=TEAL)
bar_labels(ax[0], lambda v: f"{v:.2f}%")
ax[0].set(title="(a) Median like ratio by title length", xlabel="Title length (characters)", ylabel="Likes per 100 views (%)")
ax[1].bar(tg.index.astype(str), tg.median_views, color=NAVY)
bar_labels(ax[1], kfmt)
ax[1].yaxis.set_major_formatter(M); ax[1].set(title="(b) Median views by number of tags", xlabel="Number of tags", ylabel="Median views")
save(fig, "08_title_and_tag_effects.png", 8, "Shorter titles earn a higher like ratio; more tags show higher median views",
     "Associations only, not proof of cause: tag count and title length have near-zero rank correlation with views "
     f"({corr.loc['views','tag_count']:.2f} and {corr.loc['views','title_length']:.2f}).")

# ============================================================ 7. Channels & top videos
top_ch = df.groupby("channel_title").agg(video_days=("video_id", "size"), unique_videos=("video_id", "nunique"),
                                        best_views=("views", "max")).sort_values("video_days", ascending=False).head(15)
top_ch.to_csv(RP / "top_channels.csv")
top_vid = (df.sort_values("views", ascending=False).drop_duplicates("video_id").head(10)
           [["title", "channel_title", "category_name", "views", "likes", "dislikes"]])
top_vid.to_csv(RP / "top_videos_by_views.csv", index=False)
fig, ax = plt.subplots(figsize=(11, 6.4))
t = top_ch.head(10).sort_values("video_days")
ax.barh(t.index, t.video_days, color=NAVY)
for i, (v, u) in enumerate(zip(t.video_days, t.unique_videos)):
    ax.text(v + 2, i, f"{v} video-days ({u} videos)", va="center", fontsize=8.5)
ax.margins(x=0.28)
ax.set(title="Top 10 channels by days on the trending list", xlabel="Video-days on trending list", ylabel="Channel")
save(fig, "09_top_channels.png", 9, "Late-night shows, sports and media brands appear most often",
     f"The top 10 channels account for only {df.channel_title.value_counts().head(10).sum()/len(df)*100:.1f}% of all rows, "
     "so the list is spread across 2,203 channels.")

# ============================================================ 8. Anomalies
old = df[df.days_to_trend > 365]
ratio = df.dropna(subset=["like_dislike_ratio"]).query("views > 100000")
controversial = ratio[ratio.like_dislike_ratio < 0.5].drop_duplicates("video_id")
controversial[["title", "channel_title", "views", "likes", "dislikes", "like_dislike_ratio"]].sort_values(
    "like_dislike_ratio").to_csv(RP / "anomaly_controversial_videos.csv", index=False)
old.sort_values("days_to_trend", ascending=False).drop_duplicates("video_id")[
    ["title", "channel_title", "publish_time", "trending_date", "days_to_trend", "views"]].to_csv(
    RP / "anomaly_old_videos_revived.csv", index=False)
df[df.views > iqr_hi].sort_values("views", ascending=False).drop_duplicates("video_id").head(50)[
    ["title", "channel_title", "category_name", "views"]].to_csv(RP / "anomaly_top_view_outliers.csv", index=False)

fig, ax = plt.subplots(1, 3, figsize=(18, 6.2))
ax[0].hist(np.log10(df.days_to_trend.clip(lower=0.1)), bins=50, color=NAVY, edgecolor="white", linewidth=.3)
ax[0].axvline(np.log10(365), color=ORANGE, ls="--", lw=2.2, label=f"365 days ({len(old)} rows beyond)")
ax[0].set_xticks([-1, 0, 1, 2, 3]); ax[0].set_xticklabels(["0.1", "1", "10", "100", "1,000"])
ax[0].set(title="(a) Days from upload to trending", xlabel="Days (log scale)", ylabel="Video-days"); ax[0].set_ylim(0, ax[0].get_ylim()[1] * 1.22); ax[0].legend(loc="upper right", fontsize=9)
r = df.dropna(subset=["likes", "dislikes"]); r = r[(r.likes > 0) & (r.dislikes > 0)]
ax[1].scatter(r.likes, r.dislikes, s=4, alpha=.25, color=GREY, label="All videos")
ax[1].scatter(controversial.likes, controversial.dislikes, s=28, color=ORANGE,
              label=f"More dislikes than likes, 100K+ views ({len(controversial)})")
ax[1].set_xscale("log"); ax[1].set_yscale("log"); ax[1].legend(fontsize=8, loc="upper left")
ax[1].set(title="(b) Controversial videos: likes vs dislikes", xlabel="Likes (log scale)", ylabel="Dislikes (log scale)")
dis = df.groupby("category_name").comments_disabled.mean().mul(100).sort_values()
ax[2].barh(dis.index, dis.values, color=TEAL)
for i, v in enumerate(dis.values):
    ax[2].text(v + .08, i, f"{v:.1f}%", va="center", fontsize=8)
ax[2].margins(x=0.15)
ax[2].set(title="(c) Share of video-days with comments disabled", xlabel="% of video-days", ylabel="Category")
save(fig, "10_anomalies.png", 10, "Anomalies: revived old videos, controversial videos and disabled comments",
     f"(a) Longest gap: 4,214.6 days (about 11.5 years, a Budweiser ad). (b) Includes YouTube Rewind 2017 with "
     f"1.64M dislikes. (c) Comments are disabled most often in News & Politics and Nonprofits & Activism (about 7%).")

# ============================================================ 9. Growth during trending run
g = df.sort_values("trending_date").groupby("video_id").agg(first=("views", "first"), last=("views", "last"), n=("views", "size"))
g = g[g.n >= 5]
growth = (g["last"] / g["first"])
fig, ax = plt.subplots(figsize=(10, 6))
ax.hist(growth.clip(upper=8), bins=40, color=TEAL, edgecolor="white", linewidth=.3)
ax.axvline(growth.median(), color=ORANGE, ls="--", lw=2.2, label=f"Median = {growth.median():.1f}x")
ax.set(title="Growth multiple of views for videos with 5+ trending days",
       xlabel="Views on last trending day / views on first trending day (capped at 8x)", ylabel="Number of videos")
ax.legend()
save(fig, "11_views_growth_during_trending.png", 11, "Views roughly double while a video stays on the list",
     f"Based on {len(growth):,} videos that stayed on the list 5 or more days. Median stay for all videos = 6 days "
     "(maximum 29 days). Views are cumulative snapshots.")

PDF.close()

# ============================================================ 10. Key numbers for the README + Excel tables
within1 = (df.days_to_trend <= 1).mean() * 100
facts = {
    "views_median": df.views.median(), "views_mean": df.views.mean(), "views_max": df.views.max(),
    "iqr_outlier_rows": n_out, "iqr_outlier_pct": n_out / len(df) * 100, "iqr_upper_bound": iqr_hi,
    "weekend_share_of_rows_pct": wk_share, "weekend_median_views": wk_med[True], "weekday_median_views": wk_med[False],
    "median_days_to_trend": df.days_to_trend.median(), "trend_within_1_day_pct": within1,
    "median_growth_multiple": growth.median(), "videos_in_growth_calc": len(growth),
    "old_video_rows": len(old), "controversial_videos": len(controversial),
    "median_days_on_list": vid.days_on_list.median(), "max_days_on_list": vid.days_on_list.max(),
    "videos_10plus_days": int((vid.days_on_list >= 10).sum()),
    "peak_hours_share_14_17_pct": df.publish_hour.between(14, 17).mean() * 100,
    "spearman_views_likes": corr.loc["views", "likes"], "spearman_views_comments": corr.loc["views", "comment_count"],
    "spearman_views_tags": corr.loc["views", "tag_count"], "spearman_views_title_len": corr.loc["views", "title_length"],
    "top10_channels_share_pct": df.channel_title.value_counts().head(10).sum() / len(df) * 100,
}
pd.Series(facts).round(3).to_csv(RP / "key_numbers.csv", header=["value"])

with pd.ExcelWriter(RP / "EDA_summary_tables.xlsx", engine="openpyxl") as xw:
    for name, t in [("Overview", pd.Series(overview).rename("value").to_frame()), ("Descriptive_Stats", desc.round(2)),
                    ("Category_Summary", cat.round(4)), ("Publish_Day", day_tbl), ("Monthly_Trend", mon.round(2)),
                    ("Correlation", corr.round(3)), ("Top_Channels", top_ch), ("Top_Videos", top_vid.set_index("title")),
                    ("Key_Numbers", pd.Series(facts).round(3).rename("value").to_frame())]:
        t.to_excel(xw, sheet_name=name)
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    for ws in xw.book.worksheets:
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
        ws.freeze_panes = "B2"
        for i, col in enumerate(ws.iter_cols(), 1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(c.value or "")) for c in col) + 2, 55)

for k, v in facts.items():
    print(f"{k:32s} {v:,.3f}" if isinstance(v, float) else f"{k:32s} {v}")
